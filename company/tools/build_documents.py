"""Build editable Word preparation documents without completed personal information.

Use the Codex bundled Python runtime with python-docx installed. Final DOCX files,
private inputs and QA output are deliberately excluded from Git.
"""

import argparse
import importlib.util
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("pack_content", ROOT / "content" / "registration_pack.py")
content = importlib.util.module_from_spec(spec)
spec.loader.exec_module(content)


def font(style, size, bold=False):
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(5)
    style.paragraph_format.line_spacing = 1.06


def configure(doc, title):
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(0.65), Inches(0.65)
    section.left_margin, section.right_margin = Inches(0.85), Inches(0.85)
    section.footer_distance = Inches(0.3)
    font(doc.styles["Normal"], 11)
    font(doc.styles["Title"], 19, True)
    font(doc.styles["Subtitle"], 10)
    font(doc.styles["Heading 1"], 13, True)
    font(doc.styles["Heading 2"], 11, True)
    font(doc.styles["List Bullet"], 11)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(9)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(7)
    doc.styles["Title"].paragraph_format.space_after = Pt(6)
    for name in ["Title", "Subtitle", "Heading 1", "Heading 2"]:
        style = doc.styles[name]
        style.paragraph_format.keep_with_next = True
        for border in style.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
        for color in style.element.xpath(".//w:color"):
            for key in ["themeColor", "themeTint", "themeShade"]:
                color.attrib.pop(qn("w:" + key), None)
            color.set(qn("w:val"), "000000")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("Goal Hint  |  ")
    r.font.name, r.font.size = "Arial", Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = title
    doc.core_properties.subject = "Goal Hint company registration and merchant preparation"
    doc.core_properties.author = "Goal Hint"
    doc.core_properties.keywords = "Uganda, SMC, company registration, merchant collections"


def hyperlink(paragraph, text, url):
    link = OxmlElement("w:hyperlink")
    rid = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    value = OxmlElement("w:t")
    value.text = text
    run.append(value)
    link.append(run)
    paragraph._p.append(link)


def add_table(doc, block):
    t = doc.add_table(rows=1, cols=len(block["headers"]))
    t.autofit = False
    widths = block.get("widths") or [6.8 / len(block["headers"])] * len(block["headers"])
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for cell, value in zip(t.rows[0].cells, block["headers"]):
        cell.text = value
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "E7E6E6")
        cell._tc.get_or_add_tcPr().append(shade)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in block["rows"]:
        for cell, value in zip(t.add_row().cells, values):
            cell.text = value
    borders = OxmlElement("w:tblBorders")
    for name in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        border = OxmlElement("w:" + name)
        for key, value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
            border.set(qn("w:" + key), value)
        borders.append(border)
    t._tbl.tblPr.append(borders)
    for row in t.rows:
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side, value in [("top", 85), ("bottom", 85), ("left", 100), ("right", 100)]:
                element = OxmlElement("w:" + side)
                element.set(qn("w:w"), str(value))
                element.set(qn("w:type"), "dxa")
                margins.append(element)
            props.append(margins)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.font.color.rgb = RGBColor(0, 0, 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_block(doc, block):
    kind = block["type"]
    if kind == "paragraph":
        doc.add_paragraph(block["text"])
    elif kind == "heading":
        doc.add_heading(block["text"], level=2)
    elif kind == "bullets":
        for item in block["items"]:
            doc.add_paragraph(item, style="List Bullet")
    elif kind == "fields":
        for label, value in block["items"]:
            para = doc.add_paragraph()
            para.add_run(label + ": ").bold = True
            para.add_run(value)
            para.paragraph_format.space_after = Pt(7)
    elif kind == "table":
        add_table(doc, block)
    elif kind == "reference":
        para = doc.add_paragraph()
        hyperlink(para, block["title"], block["url"])
        para.paragraph_format.space_after = Pt(7)
    else:
        raise ValueError(f"Unknown block: {kind}")


def policy_documents():
    source = ROOT / "content" / "customer-policies.json"
    raw = json.loads(source.read_text(encoding="utf-8"))
    if isinstance(raw, dict):
        raw = raw.get("documents", [])
    result = []
    for item in raw:
        pages = []
        # Explicit page groups are allowed; otherwise paginate by section budget.
        for group in item.get("pages", []):
            blocks = []
            for section in group["sections"]:
                blocks.append(content.h(section["heading"]))
                blocks += [content.p(x) for x in section.get("paragraphs", [])]
                if section.get("bullets"):
                    blocks.append(content.bullets(*section["bullets"]))
            pages.append(content.page(group.get("heading", ""), *blocks))
        if not pages:
            blocks = []
            words = 0
            for section in item["sections"]:
                count = len(" ".join(section.get("paragraphs", []) + section.get("bullets", [])).split()) + 25
                if words + count > 500 and blocks:
                    pages.append(content.page("", *blocks))
                    blocks, words = [], 0
                blocks.append(content.h(section["heading"]))
                blocks += [content.p(x) for x in section.get("paragraphs", [])]
                if section.get("bullets"):
                    blocks.append(content.bullets(*section["bullets"]))
                words += count
            if blocks:
                pages.append(content.page("", *blocks))
        name = "09-service-terms-and-refunds.docx" if "term" in item["title"].lower() else "10-privacy-notice.docx"
        result.append(content.document("03-customer-policies/" + name, item["title"], item.get("subtitle", "For approval before publication"), *pages))
    return result


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--only", nargs="*", help="Rebuild only these document stems")
    args = parser.parse_args()
    documents = content.DOCUMENTS + policy_documents()
    manifest = []
    for item in documents:
        path = ROOT / item["filename"]
        if args.only and path.stem not in args.only:
            manifest.append({"path": item["filename"], "planned_pages": len(item["pages"]), "bytes": path.stat().st_size})
            continue
        doc = Document()
        configure(doc, item["title"])
        for index, page in enumerate(item["pages"]):
            if index:
                doc.add_page_break()
            else:
                doc.add_paragraph(item["title"], style="Title")
                doc.add_paragraph(item["status"], style="Subtitle")
            if page["heading"]:
                doc.add_heading(page["heading"], level=1)
            for block in page["blocks"]:
                add_block(doc, block)
        path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(path)
        manifest.append({"path": item["filename"], "planned_pages": len(item["pages"]), "bytes": path.stat().st_size})
        print(f"Created {item['filename']} ({len(item['pages'])} planned pages)")
    work = ROOT / ".work"
    work.mkdir(exist_ok=True)
    (work / "document-manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
